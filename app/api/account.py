import io
import os
import secrets
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from flask import Blueprint, Response, g
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.database import watchlist_repository as watchlist_repo
from app.database.connection import create_connection
from app.extensions import limiter
from app.security.plans import effective_plan
from app.security.sessions import login_required

account_bp = Blueprint("account", __name__)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LOGO_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "logo.png"
)

BRAND = colors.HexColor("#5C3AED")
SOFT = colors.HexColor("#F3F0FF")
INK = colors.HexColor("#111827")
GREY = colors.HexColor("#6B7280")
LINE = colors.HexColor("#E5E7EB")

HEADER_FILL = PatternFill("solid", fgColor="5C3AED")
HEADER_FONT = Font(bold=True, color="FFFFFF")

MAX_PDF_WATCHLIST_ROWS = 40


# ---------- shared data ----------

def _load_account(user_id):
    connection = create_connection()

    try:
        cursor = connection.cursor(dictionary=True)

        cursor.execute(
            "SELECT name, email, email_verified, plan, plan_expires_at, alerts_email "
            "FROM users WHERE id = %s",
            (user_id,),
        )
        profile = cursor.fetchone() or {}

        cursor.execute(
            "SELECT plan_interval, amount_kobo, paid_at FROM payments "
            "WHERE user_id = %s AND status = 'success' ORDER BY id",
            (user_id,),
        )
        payments = cursor.fetchall()
    finally:
        connection.close()

    return profile, payments, watchlist_repo.get_watchlist_products(user_id)


def _date(value):
    if value is None:
        return ""

    return value.strftime("%d %b %Y") if hasattr(value, "strftime") else str(value)


def _number(value):
    return float(value) if value is not None else None


def _naira(value):
    return f"NGN {float(value):,.0f}" if value is not None else "-"


# ---------- Excel ----------

def _add_row(sheet, values):
    sheet.append(values)

    # Keep text cells as text so a product name starting with "="
    # can never run as a spreadsheet formula
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _style(sheet, widths):
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"


@account_bp.route("/account/export", methods=["GET"])
@limiter.limit("5 per hour")
@login_required
def export_account():
    profile, payments, watchlist_items = _load_account(g.session["id"])
    plan = effective_plan(profile)

    workbook = Workbook()

    account = workbook.active
    account.title = "Account"
    account.append(["Field", "Value"])

    for field, value in [
        ("Name", profile.get("name")),
        ("Email", profile.get("email")),
        ("Email verified", "Yes" if profile.get("email_verified") else "No"),
        ("Plan", "Pro" if plan == "pro" else "Free"),
        ("Pro active until", _date(profile.get("plan_expires_at")) if plan == "pro" else ""),
        ("Email alerts", "On" if profile.get("alerts_email") else "Off"),
        ("Theme", g.session.get("theme")),
        ("Auto refresh (minutes)", g.session.get("auto_refresh_minutes")),
        ("Default store", g.session.get("default_store")),
        ("Exported on", datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")),
    ]:
        _add_row(account, [field, value if value is not None else ""])

    _style(account, [26, 42])

    watchlist = workbook.create_sheet("Watchlist")
    watchlist.append(["Product", "Store", "Current price (₦)", "Target price (₦)", "Link", "Added on"])

    for item in watchlist_items:
        _add_row(watchlist, [
            item.get("product_name"),
            item.get("store_name"),
            _number(item.get("current_price")),
            _number(item.get("target_price")),
            item.get("product_url"),
            _date(item.get("added_at")),
        ])

        row = watchlist.max_row
        watchlist.cell(row=row, column=3).number_format = "#,##0"
        watchlist.cell(row=row, column=4).number_format = "#,##0"

    _style(watchlist, [52, 12, 18, 18, 60, 14])

    payments_sheet = workbook.create_sheet("Payments")
    payments_sheet.append(["Date", "Plan", "Amount (₦)"])

    for row in payments:
        _add_row(payments_sheet, [
            _date(row["paid_at"]),
            row["plan_interval"].capitalize(),
            row["amount_kobo"] / 100,
        ])
        payments_sheet.cell(row=payments_sheet.max_row, column=3).number_format = "#,##0"

    _style(payments_sheet, [16, 12, 16])

    buffer = io.BytesIO()
    workbook.save(buffer)

    return Response(
        buffer.getvalue(),
        mimetype=XLSX,
        headers={
            "Content-Disposition": 'attachment; filename="pricepounce-my-data.xlsx"',
            "Cache-Control": "no-store",
        },
    )


# ---------- PDF statement ----------

def _p(text, size=9, color=INK, bold=False, align=0):
    style = ParagraphStyle(
        "cell",
        fontName="Helvetica-Bold" if bold else "Helvetica",
        fontSize=size,
        leading=size + 3,
        textColor=color,
        alignment=align,
    )
    return Paragraph(escape("" if text is None else str(text)), style)


def _logo_or_wordmark():
    if os.path.exists(LOGO_PATH):
        try:
            width, height = ImageReader(LOGO_PATH).getSize()
            target_height = 30 * mm
            return Image(LOGO_PATH, width=target_height * width / height, height=target_height)
        except Exception:
            pass

    return _p("PricePounce", size=20, color=BRAND, bold=True)


def _footer(canvas, document):
    canvas.saveState()
    canvas.setStrokeColor(LINE)
    canvas.line(15 * mm, 16 * mm, A4[0] - 15 * mm, 16 * mm)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(GREY)
    canvas.drawString(
        15 * mm, 11 * mm,
        "Generated by PricePounce. This statement is private to your account.",
    )
    canvas.drawRightString(A4[0] - 15 * mm, 11 * mm, f"Page {document.page}")
    canvas.restoreState()


def _table(rows, widths, right_columns=()):
    table = Table(rows, colWidths=widths, repeatRows=1)

    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINE),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, SOFT]),
    ]

    for column in right_columns:
        commands.append(("ALIGN", (column, 0), (column, -1), "RIGHT"))

    table.setStyle(TableStyle(commands))
    return table


@account_bp.route("/account/export.pdf", methods=["GET"])
@limiter.limit("5 per hour")
@login_required
def export_account_pdf():
    profile, payments, watchlist_items = _load_account(g.session["id"])
    plan = effective_plan(profile)
    now = datetime.now(timezone.utc)

    total_paid = sum(row["amount_kobo"] for row in payments) / 100

    story = []

    # Header
    header = Table(
        [[_logo_or_wordmark(), [
            _p("ACCOUNT STATEMENT", size=13, color=BRAND, bold=True, align=2),
            _p(f"No. PP-{secrets.token_hex(4).upper()}", size=8, color=GREY, align=2),
        ]]],
        colWidths=[90 * mm, 90 * mm],
    )
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, 0), 2, BRAND),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    story += [header, Spacer(1, 8 * mm)]

    # Who and when
    details = Table(
        [[
            [
                _p("PREPARED FOR", size=7, color=GREY, bold=True),
                _p(profile.get("name") or "-", size=11, bold=True),
                _p(profile.get("email") or "-", size=9, color=GREY),
            ],
            [
                _p("DATE ISSUED", size=7, color=GREY, bold=True, align=2),
                _p(now.strftime("%d %B %Y"), size=11, bold=True, align=2),
                _p(now.strftime("%H:%M UTC"), size=9, color=GREY, align=2),
            ],
        ]],
        colWidths=[90 * mm, 90 * mm],
    )
    details.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story += [details, Spacer(1, 6 * mm)]

    # Summary tiles
    plan_text = "Pro" if plan == "pro" else "Free"

    if plan == "pro" and profile.get("plan_expires_at"):
        plan_note = f"Active until {_date(profile['plan_expires_at'])}"
    else:
        plan_note = "Upgrade anytime in the app"

    def tile(label, value, note):
        return [
            _p(label, size=7, color=GREY, bold=True),
            _p(value, size=15, color=BRAND, bold=True),
            _p(note, size=8, color=GREY),
        ]

    summary = Table(
        [[
            tile("PLAN", plan_text, plan_note),
            tile("WATCHLIST", str(len(watchlist_items)), "products tracked"),
            tile("TOTAL PAID", _naira(total_paid), f"{len(payments)} payment(s)"),
        ]],
        colWidths=[60 * mm, 60 * mm, 60 * mm],
    )
    summary.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), SOFT),
        ("BOX", (0, 0), (0, 0), 3, colors.white),
        ("BOX", (1, 0), (1, 0), 3, colors.white),
        ("BOX", (2, 0), (2, 0), 3, colors.white),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    story += [summary, Spacer(1, 8 * mm)]

    # Watchlist
    story += [_p("Your watchlist", size=12, bold=True), Spacer(1, 3 * mm)]

    if watchlist_items:
        rows = [[
            _p("Product", color=colors.white, bold=True),
            _p("Store", color=colors.white, bold=True),
            _p("Price", color=colors.white, bold=True, align=2),
            _p("Target", color=colors.white, bold=True, align=2),
        ]]

        for item in watchlist_items[:MAX_PDF_WATCHLIST_ROWS]:
            rows.append([
                _p(item.get("product_name"), size=8.5),
                _p(item.get("store_name"), size=8.5),
                _p(_naira(item.get("current_price")), size=8.5, align=2),
                _p(_naira(item.get("target_price")), size=8.5, align=2),
            ])

        story.append(_table(rows, [98 * mm, 20 * mm, 31 * mm, 31 * mm], right_columns=(2, 3)))

        extra = len(watchlist_items) - MAX_PDF_WATCHLIST_ROWS

        if extra > 0:
            story += [
                Spacer(1, 2 * mm),
                _p(f"and {extra} more. Use the Excel export for the full list.", size=8, color=GREY),
            ]
    else:
        story.append(_p("You aren't tracking any products yet.", size=9, color=GREY))

    story.append(Spacer(1, 8 * mm))

    # Payments
    story += [_p("Payment history", size=12, bold=True), Spacer(1, 3 * mm)]

    if payments:
        rows = [[
            _p("Date", color=colors.white, bold=True),
            _p("Plan", color=colors.white, bold=True),
            _p("Amount", color=colors.white, bold=True, align=2),
        ]]

        for row in payments:
            rows.append([
                _p(_date(row["paid_at"]), size=8.5),
                _p(row["plan_interval"].capitalize(), size=8.5),
                _p(_naira(row["amount_kobo"] / 100), size=8.5, align=2),
            ])

        story.append(_table(rows, [60 * mm, 60 * mm, 60 * mm], right_columns=(2,)))
    else:
        story.append(_p("No payments yet.", size=9, color=GREY))

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=22 * mm,
        title="PricePounce account statement",
        author="PricePounce",
    )
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)

    return Response(
        buffer.getvalue(),
        mimetype="application/pdf",
        headers={
            "Content-Disposition": 'attachment; filename="pricepounce-statement.pdf"',
            "Cache-Control": "no-store",
        },
    )